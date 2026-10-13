#!/usr/bin/env python3
"""Apply one explicit JSON patch; preserve untouched DOCX package parts."""
import argparse
import copy
import hashlib
import io
import json
import os
import re
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as ET
from inspect_report import (W, NS, FILENAME_TOKEN, SHOT, has_image, text_of,
                            label_of, read_package, region_nodes, resolve_input)

PROFILE_FIELDS = {'name', 'student_id', 'class', 'major', 'college', 'school'}
AUTO_LABEL = re.compile(r'建议评语|建议评分|供(?:教师|老师)(?:参考|审核)|待(?:教师|老师)确认|自动(?:生成|拟写)|AI\s*(?:拟写|生成)|建议稿|[（(]\s*建议\s*[）)]|(?:^|\n)\s*建议\s*[:：]', re.I)


def validate_generated_text(value, field=None):
    # Check new material only. Normal prose containing 建议 and existing reviews stay intact.
    if AUTO_LABEL.search(value) or value.strip() == '建议' or (field == 'score' and '建议' in value):
        raise ValueError('Remove automatic review/disclaimer labels from new text')
    return value


def read_profile(path):
    if not path.is_file():
        return {}
    values = {}
    for line in path.read_text(encoding='utf-8-sig').splitlines():
        match = re.fullmatch(r'(name|student_id|class|major|college|school):\s*(.*?)\s*', line)
        if not match:
            continue
        key, value = match.groups()
        if value.startswith('"'):
            value = json.loads(value)
        elif value.startswith("'") and value.endswith("'"):
            value = value[1:-1].replace("''", "'")
        elif value.startswith(('|', '>')):
            raise ValueError('Profile fields must be single-line strings')
        values[key] = str(value).strip()
    return values


def need(profile, key):
    value = profile.get(key, '')
    if not value:
        raise ValueError('Missing student-profile.md field: ' + key)
    return value


def last_two(profile):
    value = need(profile, 'student_id')
    if len(value) < 2:
        raise ValueError('student_id must contain at least 2 characters')
    return value[-2:]


def output_name(source, profile):
    def replace(m):
        token = m.group()
        if '最后' in token:
            return last_two(profile)
        key = {'学号': 'student_id', '姓名': 'name', '班级': 'class', '专业班级': 'class'}[token]
        return need(profile, key)
    name = FILENAME_TOKEN.sub(replace, source.name)
    if name == source.name:
        name = re.sub(r'(?:\s*[（(]\d+[）)])+(?=\.docx$)', '', name, flags=re.I)
    name = name.replace('+', '')
    if name == source.name:
        name = source.stem + '_已填写' + source.suffix
    if re.search(r'[<>:"/\\|?*\x00-\x1f]', name) or name.rstrip(' .') != name:
        raise ValueError('Invalid output filename; check identity fields or supply an explicit output path')
    return name


def signature_region(node):
    # A paragraph target must not bypass a signature in the same teacher cell.
    for candidate in [node, *node.iterancestors()]:
        if candidate.tag in (W + 'tc', W + 'p') and '指导教师' in label_of(candidate) and has_image(candidate):
            return True
    return False


def insert_image(node, op, image_doc, part_name, patch_dir, drawing_id):
    from docx.oxml import OxmlElement
    from docx.shared import Inches
    from docx.text.paragraph import Paragraph
    if node.tag != W + 'p':
        raise ValueError('insert_image requires an inspected paragraph ID')
    if text_of(node) != op.get('expect', ''):
        raise ValueError('insert_image precondition failed')
    evidence = op.get('evidence', {})
    if evidence.get('verified') is not True or not evidence.get('source') or not evidence.get('requirement'):
        raise ValueError('Image requires agent-verified source and requirement evidence')
    width = op.get('width_inches')
    if not isinstance(width, (float, int)) or isinstance(width, bool) or width <= 0:
        raise ValueError('Supply width_inches based on the available template width')
    picture = Path(op['image'])
    if not picture.is_absolute():
        picture = patch_dir / picture
    part = next(p for p in image_doc.part.package.parts if str(p.partname) == '/' + part_name)
    paragraph = Paragraph(OxmlElement('w:p'), part)
    paragraph.add_run().add_picture(str(picture), width=Inches(width))
    marker = op.get('marker')
    if marker:
        if not SHOT.fullmatch(marker):
            raise ValueError('marker must be an actual screenshot prompt')
        replace_text(node, marker, '')
    elif SHOT.search(text_of(node)):
        raise ValueError('Specify the screenshot marker fulfilled by this image')
    drawing = ET.fromstring(paragraph.runs[0]._r.xml.encode('utf-8'))
    for prop in drawing.xpath('.//*[local-name()="docPr"]'):
        prop.set('id', str(drawing_id))
    node.append(drawing)
    caption = op.get('caption')
    if caption:
        validate_generated_text(caption)
        p = ET.Element(W + 'p')
        props = node.find(W + 'pPr')
        if props is not None:
            p.append(copy.deepcopy(props))
        style = next(iter(node.xpath('./w:r/w:rPr', namespaces=NS)), None)
        r = ET.SubElement(p, W + 'r')
        if style is not None:
            r.append(copy.deepcopy(style))
        ET.SubElement(r, W + 't').text = caption
        node.addnext(p)


def date_value(op, patch, existing):
    # Agent supplies explicit report/manual/user date and template format in the patch.
    selected = op.get('value') or patch.get('date')
    if not selected and re.search(r'\d{4}\D+\d{1,2}\D+\d{1,2}', existing):
        return existing  # An existing explicit field is not overwritten by today's date.
    selected = selected or patch.get('local_date') or datetime.now().astimezone().date().isoformat()
    digits = re.fullmatch(r'(\d{4})\D+(\d{1,2})\D+(\d{1,2})\D*', str(selected))
    if not digits:
        raise ValueError('Date must contain year, month and day')
    d = datetime(*map(int, digits.groups()))
    fmt = op.get('format') or patch.get('date_format')
    if not fmt:
        fmt = 'YYYY年MM月DD日' if '年' in existing and '月' in existing else 'YYYY.MM.DD'
        if '/' in existing:
            fmt = 'YYYY/MM/DD'
        elif '-' in existing:
            fmt = 'YYYY-MM-DD'
    if fmt not in ('YYYY.MM.DD', 'YYYY-MM-DD', 'YYYY/MM/DD', 'YYYY年MM月DD日'):
        raise ValueError('Unsupported date format; use a text operation for other explicit template formats')
    return fmt.replace('YYYY', f'{d.year:04}').replace('MM', f'{d.month:02}').replace('DD', f'{d.day:02}')


def put_text(node, value, clean_prompt=False, font_size_pt=None):
    if font_size_pt is not None:
        if (isinstance(font_size_pt, bool) or not isinstance(font_size_pt, (int, float))
                or not 0 < font_size_pt <= 1638 or not float(font_size_pt * 2).is_integer()):
            raise ValueError("font_size_pt must be a positive half-point value up to 1638")
    if node.xpath('.//w:tbl | .//w:txbxContent | .//w:fldChar | .//w:sdt | .//w:hyperlink', namespaces=NS):
        raise ValueError('Complex region: use a precise replace operation or existing document tools')
    if has_image(node):
        raise ValueError('set_text cannot replace a region containing an image')
    paragraphs = [node] if node.tag == W + 'p' else node.findall(W + 'p')
    if not paragraphs:
        raise ValueError('No editable paragraph')
    first = paragraphs[0]
    style = next(iter(first.xpath('./w:r/w:rPr', namespaces=NS)), None)
    for t in node.xpath('.//w:t | .//w:br | .//w:tab', namespaces=NS):
        t.getparent().remove(t)
    run = ET.SubElement(first, W + 'r')
    if style is not None:
        run.append(copy.deepcopy(style))
    if font_size_pt is not None:
        props = run.find(W + 'rPr')
        if props is None:
            props = ET.SubElement(run, W + 'rPr')
        for tag in ('sz', 'szCs'):
            sizes = props.findall(W + tag)
            size = sizes[0] if sizes else ET.SubElement(props, W + tag)
            for duplicate in sizes[1:]:
                props.remove(duplicate)
            size.set(W + 'val', str(int(font_size_pt * 2)))
    for index, line in enumerate(str(value).split('\n')):
        if index:
            ET.SubElement(run, W + 'br')
        t = ET.SubElement(run, W + 't')
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        t.text = line
    if clean_prompt:
        for prop in node.xpath('.//w:highlight | .//w:color | ./w:tcPr/w:shd | .//w:pPr/w:shd', namespaces=NS):
            prop.getparent().remove(prop)


def indent_report_paragraphs(node):
    if node.tag not in (W + 'p', W + 'tc') or node.xpath('.//w:tbl | .//w:txbxContent', namespaces=NS):
        raise ValueError('Select a prose paragraph or leaf cell for indentation')
    paragraphs = [node] if node.tag == W + 'p' else node.findall(W + 'p')
    for p in paragraphs:
        if not text_of(p).strip() or has_image(p) or p.find('./' + W + 'pPr/' + W + 'numPr') is not None:
            continue
        props = p.find(W + 'pPr')
        if props is None:
            props = ET.Element(W + 'pPr')
            p.insert(0, props)
        ind = props.find(W + 'ind')
        if ind is None:
            ind = ET.SubElement(props, W + 'ind')
        for attr in ('firstLine', 'hanging', 'hangingChars'):
            ind.attrib.pop(W + attr, None)
        ind.set(W + 'firstLineChars', '200')


def put_report_prose(node, value, clean_prompt=False, font_size_pt=None):
    # Reuse the guarded text writer; each newline in these two prose fields is a paragraph.
    paragraphs = [node] if node.tag == W + 'p' else node.findall(W + 'p')
    if not paragraphs:
        raise ValueError('No editable paragraph')
    saved_props = copy.deepcopy(paragraphs[0].find(W + 'pPr'))
    saved_style = next(iter(paragraphs[0].xpath('./w:r/w:rPr', namespaces=NS)), None)
    saved_style = copy.deepcopy(saved_style)
    lines = str(value).split('\n')
    put_text(node, lines[0], clean_prompt, font_size_pt)
    for index, line in enumerate(lines[1:], 1):
        if index < len(paragraphs):
            p = paragraphs[index]
        else:
            p = ET.Element(W + 'p')
            if saved_props is not None:
                p.append(copy.deepcopy(saved_props))
            paragraphs[-1].addnext(p)
            paragraphs.append(p)
        if saved_style is not None and not p.xpath('./w:r/w:rPr', namespaces=NS):
            ET.SubElement(p, W + 'r').append(copy.deepcopy(saved_style))
        put_text(p, line, clean_prompt, font_size_pt)
    for p in paragraphs:
        indent_report_paragraphs(p)


def replace_text(node, old, new):
    texts = node.xpath('.//w:t', namespaces=NS)
    joined = ''.join(t.text or '' for t in texts)
    if not old or joined.count(old) != 1:
        raise ValueError('replace.old must match exactly once within the selected region')
    start = joined.index(old)
    end = start + len(old)
    cursor = 0
    inserted = False
    for t in texts:
        value = t.text or ''
        a, b = cursor, cursor + len(value)
        if a < end and b > start:
            prefix = value[:max(0, start - a)]
            suffix = value[max(0, end - a):]
            t.text = prefix + (new if not inserted else '') + suffix
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            inserted = True
        cursor = b


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('docx')
    parser.add_argument('patch')
    parser.add_argument('--output')
    parser.add_argument('--profile', help='Default: installed skill references/student-profile.md')
    parser.add_argument('--overwrite-generated', action='store_true', help='Only for a confirmed previous generated output')
    args = parser.parse_args()
    source = resolve_input(args.docx)
    patch = json.loads(Path(args.patch).read_text(encoding='utf-8-sig'))
    raw, entries, trees = read_package(source)
    if patch.get('source_sha256') != hashlib.sha256(raw).hexdigest():
        raise ValueError('Missing/stale source_sha256: use the current inspection snapshot')
    profile_path = Path(args.profile) if args.profile else Path(__file__).resolve().parents[1] / 'references/student-profile.md'
    profile = read_profile(profile_path)
    output = Path(args.output).resolve() if args.output else source.with_name(output_name(source, profile))
    output = output.with_name(output.name.replace('+', ''))
    if output.suffix.lower() != '.docx':
        raise ValueError('Output must use the .docx extension')
    if output == source or (output.exists() and os.path.samefile(output, source)):
        raise ValueError('Never overwrite the original input')
    if output.exists() and not args.overwrite_generated:
        raise ValueError('Output exists; confirm it is a previous generated output before overwriting')
    nodes = region_nodes(trees)
    requirements = patch.get('requirements', [])
    if any(r.get('status') not in ('原要求完成', '提供替代成果', '仍未完成') for r in requirements):
        raise ValueError('Invalid requirement completion status')
    touched, changed, skipped, pending = set(), [], [], []
    auto_teacher_nodes = []
    from docx import Document
    image_doc = None
    image_parts, image_records = set(), []
    drawing_id = 1 + max((int(v) for tree in trees.values()
                          for v in tree.xpath('//*[local-name()="docPr"]/@id')), default=0)
    deferred_scores = []
    for op in patch.get('operations', []):
        action = op.get('op', 'set_text')
        if action not in ('set_text', 'replace', 'keep', 'insert_image', 'format_paragraphs'):
            raise ValueError('Unsupported operation')
        if 'font_size_pt' in op and action != 'set_text':
            raise ValueError('font_size_pt is supported only for set_text')
        target = op['target']
        if target in touched:
            raise ValueError('One operation per target; consolidate the patch')
        touched.add(target)
        node = nodes[target]
        existing = text_of(node)
        field = op.get('field')
        teacher = field == 'teacher' or '指导教师' in label_of(node)
        if signature_region(node) or (teacher and has_image(node)):
            skipped.append({'target': target, 'reason': 'teacher_signature_image_preserved'})
            continue
        if action == 'format_paragraphs':
            if field not in ('reflection', 'teacher_comment'):
                raise ValueError('Paragraph indentation is limited to reflection and teacher_comment')
            if 'text' in op or 'clean_prompt' in op:
                raise ValueError('format_paragraphs preserves existing text and prompt styling')
            indent_report_paragraphs(node)
            changed.append(target)
            continue
        if op.get('op', 'set_text') == 'keep':
            pending.append({'target': target, 'reason': op.get('reason', 'intentionally preserved')})
            continue
        if action == 'insert_image':
            if image_doc is None:
                image_doc = Document(io.BytesIO(raw))
            part_name = target.split('#', 1)[0]
            insert_image(node, op, image_doc, part_name, Path(args.patch).resolve().parent, drawing_id)
            drawing_id += 1
            image_parts.add(part_name)
            image_records.append({'target': target, 'evidence': op['evidence'], 'image': op['image']})
            changed.append(target)
            continue
        if SHOT.search(existing):
            skipped.append({'target': target, 'reason': 'screenshot_requirement_preserved; use insert_image'})
            pending.append({'target': target, 'reason': 'screenshot_unverified'})
            continue
        if field in PROFILE_FIELDS:
            value = need(profile, field)
        elif field == 'student_id_last2':
            value = last_two(profile)
        elif field == 'date':
            value = date_value(op, patch, existing)
        elif field == 'score' and 'text' not in op:
            if action != 'set_text' or existing != op.get('expect', '') or existing.strip():
                raise ValueError('Automatic score requires an empty field')
            deferred_scores.append((target, node, op))
            continue
        else:
            value = op.get('text')
        if not isinstance(value, str):
            raise ValueError('Operation requires text or a supported field')
        validate_generated_text(value, field)
        if op.get('op', 'set_text') == 'replace':
            replace_text(node, op['old'], value)
        else:
            if existing != op.get('expect', ''):
                raise ValueError('set_text precondition failed: ' + target)
            writer = put_report_prose if field in ('reflection', 'teacher_comment') else put_text
            writer(node, value, op.get('clean_prompt', False), op.get('font_size_pt'))
        changed.append(target)
        if teacher:
            auto_teacher_nodes.append(node)
    # Overlapping parent/child targets could invalidate earlier writes: reject before saving.
    selected = [nodes[t] for t in touched]
    if any(a is not b and a in b.iterancestors() for a in selected for b in selected):
        raise ValueError('Overlapping targets: choose either cell or paragraph, not both')
    if any(has_image(n) for n in auto_teacher_nodes):
        raise ValueError('QA failed: image and automatic teacher text overlap')
    # All explicit screenshot markers remain unverified, even when the patch does not touch them.
    shot_slots = []
    for target, node in nodes.items():
        if node.tag != W + 'p':
            continue
        shot_slots.extend(target + f'@shot{i}' for i, _ in enumerate(SHOT.finditer(text_of(node))))
    deferred_targets = {t for t, _, _ in deferred_scores}
    missing = [t for t in patch.get('required', []) if t not in deferred_targets and not text_of(nodes[t]).strip() and not has_image(nodes[t])]
    partial = bool(missing or pending or shot_slots or any(r['status'] != '原要求完成' for r in requirements))
    for target, node, op in deferred_scores:
        put_text(node, '良' if partial else '优', op.get('clean_prompt', False), op.get('font_size_pt'))
        changed.append(target)
    extra = {}
    if image_doc is not None:
        buffer = io.BytesIO()
        image_doc.save(buffer)
        original_names = {info.filename for info, _ in entries}
        rel_names = {str(Path(p).parent).replace('\\', '/') + '/_rels/' + Path(p).name + '.rels' for p in image_parts}
        with ZipFile(buffer) as z:
            extra = {name: z.read(name) for name in z.namelist()
                     if name not in original_names or name in rel_names or name == '[Content_Types].xml'}
    parts = {t.split('#', 1)[0] for t in changed}
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(suffix='.docx', dir=output.parent)
    os.close(fd)
    try:
        with ZipFile(temp, 'w') as z:
            for info, data in entries:
                if info.filename in extra:
                    data = extra.pop(info.filename)
                if info.filename in parts:
                    data = ET.tostring(trees[info.filename], xml_declaration=True, encoding='UTF-8', standalone=True)
                z.writestr(info, data)
            for name, data in extra.items():
                z.writestr(name, data)
        Document(temp)  # One basic package-open QA, not a visual layout proof.
        with ZipFile(temp) as z:
            if z.testzip() is not None:
                raise ValueError('DOCX ZIP integrity failure')
        if args.overwrite_generated:
            os.replace(temp, output)
        else:
            # Windows rename fails if a destination appeared during processing.
            if os.name == 'nt':
                os.rename(temp, output)
            else:
                os.link(temp, output)
                os.unlink(temp)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
    print(json.dumps({'output': str(output), 'changed': changed, 'skipped': skipped,
                      'written_values': [{'target': t, 'text': text_of(nodes[t])[:280]} for t in changed],
                      'missing_required': missing, 'pending': pending, 'screenshot_slots_unverified': shot_slots,
                      'inserted_images': image_records, 'requirements': requirements,
                      'opens': True, 'basic_qa': 'partial' if partial else 'pass',
                      'visual_qa': 'not_performed'}, ensure_ascii=False, separators=(',', ':')))


if __name__ == '__main__':
    try:
        main()
    except (OSError, ValueError, KeyError, ET.XMLSyntaxError) as exc:
        print(json.dumps({'error': str(exc)}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)
