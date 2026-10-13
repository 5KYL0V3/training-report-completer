"""Regression checks using temporary synthetic DOCX files only."""
import base64
import copy
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from lxml import etree as ET

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
from inspect_report import W, NS, inventory, read_package, region_nodes, text_of
from apply_report import output_name


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.picture = self.root / 'synthetic.png'
        self.picture.write_bytes(base64.b64decode(
            'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII='))
        self.screenshot = self.root / 'synthetic-evidence.png'
        self.screenshot.write_bytes(self.picture.read_bytes() + b'TEST ONLY distinct synthetic image')
        self.profile = self.root / 'TEST_ONLY_profile.md'
        self.profile.write_text('name: "测试同学"\nstudent_id: "000007"\nclass: "测试班"\n', encoding='utf-8')
        doc = Document()
        doc.add_paragraph('原有正文：建议记录检索条件。')
        doc.add_paragraph('原有批阅：供教师审核')  # Existing material is not mechanically scrubbed.
        table = doc.add_table(rows=6, cols=2)
        table.style = 'Table Grid'
        for row, label in zip(table.rows, ('过程', '教师评语', '评分', '指导教师', '截图', '保留合并区域')):
            row.cells[0].text = label
        table.cell(0, 1).add_table(rows=1, cols=2).cell(0, 0).text = '内层标签'
        table.cell(3, 1).paragraphs[0].add_run().add_picture(str(self.picture), width=Inches(.2))
        table.cell(4, 1).text = '（截图）'
        table.cell(5, 0).merge(table.cell(5, 1))
        props = table.rows[0]._tr.get_or_add_trPr()
        height = OxmlElement('w:trHeight')
        height.set(W + 'val', '600')
        height.set(W + 'hRule', 'exact')
        props.append(height)
        props.append(OxmlElement('w:cantSplit'))
        shade = OxmlElement('w:shd')
        shade.set(W + 'fill', 'FFFF00')
        table.cell(0, 1)._tc.get_or_add_tcPr().append(shade)
        doc.sections[0].header.paragraphs[0].text = '页眉（截图）'
        self.source = self.root / '学号（最后2位）+姓名+班级+实训.docx'
        doc.save(self.source)
        self.raw, self.entries, self.trees = read_package(self.source)
        self.nodes = region_nodes(self.trees)
        self.report = inventory(self.source, self.trees, self.raw, self.entries)

    def cell(self, label):
        return next(key for key, node in self.nodes.items()
                    if node.tag == W + 'tc' and node.getprevious() is not None
                    and node.getprevious().tag == W + 'tc'
                    and text_of(node.getprevious()).strip() == label)

    def paragraph(self, cell):
        return next(r['paragraph_ids'][0] for r in self.report['regions'] if r['id'] == cell)

    def run_patch(self, operations, *, output=None, requirements=None, required=None, extra=(), success=True):
        patch = self.root / 'patch.json'
        patch.write_text(json.dumps({'source_sha256': hashlib.sha256(self.raw).hexdigest(),
                                     'operations': operations, 'requirements': requirements or [],
                                     'required': required or []}, ensure_ascii=False), encoding='utf-8')
        cmd = [sys.executable, '-X', 'utf8', str(SCRIPTS / 'apply_report.py'), str(self.source), str(patch),
               '--profile', str(self.profile)]
        if output:
            cmd += ['--output', str(output)]
        result = subprocess.run(cmd + list(extra), capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0 if success else 1, result.stdout + result.stderr)
        self.assertEqual(self.source.read_bytes(), self.raw, 'Original was changed')
        return json.loads(result.stdout if success else result.stderr)

    def test_names(self):
        profile = {'name': '测试同学', 'student_id': '000007', 'class': '测试班'}
        self.assertEqual(output_name(self.source, profile), '07测试同学测试班实训.docx')
        name = '07+测试同学+测试班+实训3外文文献线索检索与利用实训.docx'
        self.assertEqual(output_name(Path(name), {}), name.replace('+', ''))
        self.assertEqual(output_name(Path('报告.docx'), {}), '报告_已填写.docx')
        result = self.run_patch([], output=self.root / 'work' / '指定+输出.docx')
        self.assertEqual(Path(result['output']).name, '指定输出.docx')
        self.assertTrue(Path(result['output']).is_file())

    def test_inventory_nested_paragraphs_and_unique_shots(self):
        self.assertEqual(len(self.report['screenshots']), 2)
        self.assertTrue(all('#p' in r['target'] for r in self.report['screenshots']))
        nested = next(t for t in self.report['tables'] if t['parent_cell'])
        leaf = nested['rows'][0][1]
        self.assertIn(self.paragraph(leaf), self.nodes)
        self.assertTrue(self.report['filename_contains_plus'])

    def image_op(self, target, expect, caption):
        return {'op': 'insert_image', 'target': target, 'expect': expect, 'marker': '（截图）',
                'image': self.screenshot.name, 'width_inches': .5, 'caption': caption,
                'evidence': {'verified': True, 'source': 'TEST ONLY synthetic fixture', 'requirement': 'test slot'}}

    def test_nested_fill_images_caption_and_structure(self):
        leaf = next(t for t in self.report['tables'] if t['parent_cell'])['rows'][0][1]
        shot = self.paragraph(self.cell('截图'))
        header = next(s['target'] for s in self.report['screenshots'] if 'header' in s['target'])
        teacher = self.paragraph(self.cell('指导教师'))
        comment, score = self.cell('教师评语'), self.cell('评分')
        result = self.run_patch([
            {'target': leaf, 'text': '步骤清晰。建议保留检索条件。'},
            {'target': comment, 'field': 'teacher_comment', 'text': '报告过程清晰，检索方法运用合理。'},
            {'target': score, 'field': 'score'},
            {'target': teacher, 'field': 'teacher', 'text': '不应插入的教师姓名'},
            self.image_op(shot, '（截图）', '图1 合成测试图'),
            self.image_op(header, '页眉（截图）', '图2 合成测试图')],
            requirements=[{'id': '截图', 'status': '原要求完成'}], required=[leaf, comment, score])
        self.assertEqual(result['basic_qa'], 'pass')
        self.assertEqual(result['screenshot_slots_unverified'], [])
        self.assertEqual(result['skipped'][0]['reason'], 'teacher_signature_image_preserved')
        out_raw, out_entries, out_trees = read_package(result['output'])
        body = text_of(out_trees['word/document.xml'])
        self.assertIn('图1 合成测试图', body)
        self.assertIn('建议保留检索条件', body)
        self.assertIn('原有批阅：供教师审核', body)
        self.assertNotIn('不应插入', body)
        self.assertIn({'target': score, 'text': '优'}, result['written_values'])
        def structure(tree):
            clone = copy.deepcopy(tree)
            for p in list(clone.iter(W + 'p')):
                p.getparent().remove(p)
            return ET.tostring(clone, method='c14n')
        self.assertEqual(structure(self.trees['word/document.xml']), structure(out_trees['word/document.xml']))
        original_images = self.trees['word/document.xml'].xpath('.//w:drawing', namespaces=NS)
        new_images = out_trees['word/document.xml'].xpath('.//w:drawing', namespaces=NS)
        self.assertEqual(ET.tostring(original_images[0]), ET.tostring(new_images[0]))
        old_parts = {i.filename: data for i, data in self.entries}
        new_parts = {i.filename: data for i, data in out_entries}
        self.assertEqual(sum(n.startswith('word/media/') for n in new_parts), 2)
        allowed = {'word/document.xml', 'word/header1.xml', 'word/_rels/document.xml.rels',
                   'word/_rels/header1.xml.rels', '[Content_Types].xml'}
        for name, data in old_parts.items():
            if name not in allowed:
                self.assertEqual(data, new_parts[name], name)
        all_ids = [v for tree in out_trees.values() for v in tree.xpath('//*[local-name()="docPr"]/@id')]
        self.assertEqual(len(all_ids), len(set(all_ids)))
        self.assertEqual(len(Document(result['output']).tables), 1)

    def test_labels_rejected_before_save(self):
        for index, text in enumerate(('良（建议）', '建议评语：完整', '供教师审核', '待教师确认', '自动生成', 'AI拟写', '建议：优')):
            with self.subTest(text=text):
                target = self.root / f'blocked-{index}.docx'
                error = self.run_patch([{'target': self.cell('教师评语'), 'text': text}], output=target, success=False)
                self.assertIn('labels', error['error'])
                self.assertFalse(target.exists())

    def test_partial_grade_and_evidence(self):
        result = self.run_patch([{'target': self.cell('评分'), 'field': 'score'},
                                 {'target': self.cell('截图'), 'expect': '（截图）', 'text': '声称已完成'}],
                                requirements=[{'id': '导出', 'status': '提供替代成果'}])
        self.assertEqual(result['basic_qa'], 'partial')
        self.assertIn({'target': self.cell('评分'), 'text': '良'}, result['written_values'])
        self.assertEqual(len(result['screenshot_slots_unverified']), 2)
        self.assertEqual(result['requirements'][0]['status'], '提供替代成果')
        bad = self.image_op(self.paragraph(self.cell('截图')), '（截图）', '图1')
        bad['evidence']['verified'] = False
        self.assertIn('agent-verified', self.run_patch([bad], output=self.root / 'bad.docx', success=False)['error'])

    def test_original_collision_and_unknown_destination(self):
        # Use a source with no plus, to test equality after output-name normalization.
        protected = self.root / '原件.docx'
        protected.write_bytes(self.raw)
        self.source = protected
        for extra in ((), ('--overwrite-generated',)):
            self.assertIn('original', self.run_patch([], output=protected, extra=extra, success=False)['error'])
        existing = self.root / '未知.docx'
        existing.write_bytes(b'unknown user file')
        self.assertIn('exists', self.run_patch([], output=self.root / '未+知.docx', success=False)['error'])
        self.assertEqual(existing.read_bytes(), b'unknown user file')
        hardlink = self.root / 'hardlink.docx'
        os.link(protected, hardlink)
        self.assertIn('original', self.run_patch([], output=hardlink, extra=('--overwrite-generated',), success=False)['error'])

    def test_external_requirement_controls_grade_without_screenshot_blockers(self):
        # Removing test prompts here builds a separate synthetic fixture, never a user file.
        doc = Document(self.source)
        doc.tables[0].cell(4, 1).text = ''
        doc.sections[0].header.paragraphs[0].text = ''
        doc.save(self.source)
        self.raw, self.entries, self.trees = read_package(self.source)
        self.nodes = region_nodes(self.trees)
        self.report = inventory(self.source, self.trees, self.raw, self.entries)
        for index, (status, grade) in enumerate((('原要求完成', '优'), ('提供替代成果', '良'), ('仍未完成', '良'))):
            result = self.run_patch([{'target': self.cell('评分'), 'field': 'score'}],
                                    output=self.root / f'status-{index}.docx',
                                    requirements=[{'id': '原生导出', 'status': status}])
            self.assertEqual(result['screenshot_slots_unverified'], [])
            self.assertIn({'target': self.cell('评分'), 'text': grade}, result['written_values'])
            self.assertEqual(result['basic_qa'], 'pass' if status == '原要求完成' else 'partial')

    def test_parent_reset_and_overlap_rejected(self):
        parent = self.cell('过程')
        error = self.run_patch([{'target': parent, 'expect': text_of(self.nodes[parent]), 'text': '重置'}], success=False)
        self.assertIn('Complex region', error['error'])
        cell = self.cell('教师评语')
        error = self.run_patch([{'target': cell, 'text': '评语'},
                                {'target': self.paragraph(cell), 'expect': '评语', 'text': '再次改写'}], success=False)
        self.assertIn('Overlapping', error['error'])


if __name__ == '__main__':
    unittest.main()
