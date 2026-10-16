"""Revision mode uses explicit targets and synthetic identity data only."""
import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import unittest
from docx import Document
from lxml import etree as ET
import test_report as fixtures
from inspect_report import W, NS, read_package, region_nodes, inventory, text_of
from apply_report import output_name


class RevisionTests(unittest.TestCase):
    def setUp(self):
        self.f = fixtures.ReportTests()
        self.f.setUp()
        self.addCleanup(self.f.doCleanups)

    def run_patch(self, patch, *, success=True, output=None):
        f=self.f
        patch={'source_sha256':hashlib.sha256(f.raw).hexdigest(),**patch}
        file=f.root/'revision-patch.json'
        file.write_text(json.dumps(patch,ensure_ascii=False),encoding='utf-8')
        cmd=[sys.executable,'-X','utf8',str(fixtures.SCRIPTS/'apply_report.py'),str(f.source),str(file),'--profile',str(f.profile)]
        if output: cmd+=['--output',str(output)]
        result=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0 if success else 1,result.stdout+result.stderr)
        self.assertEqual(f.source.read_bytes(),f.raw)
        return json.loads(result.stdout if success else result.stderr)

    def populated_fixture(self):
        f=self.f
        doc=Document(f.source)
        doc.add_paragraph('学生：旧同学甲')
        doc.add_paragraph('文献作者：旧同学甲；指导教师：教师乙')
        doc.add_paragraph('教师实际批阅：原记录保留')
        doc.sections[0].header.paragraphs[0].text='班级：旧测试班'
        doc.sections[0].footer.paragraphs[0].text='学号：000099'
        doc.tables[0].cell(0,1).tables[0].cell(0,1).text='原步骤错误内容'
        doc.tables[0].cell(1,1).text='原普通评语'
        cells=doc.tables[0].add_row().cells
        cells[0].text='心得'; cells[1].text='原心得第一段'
        cells[1].add_paragraph('原心得第二段')
        f.source=f.root/'99+旧同学甲+旧测试班+实训.docx'
        doc.save(f.source)
        f.raw,f.entries,f.trees=read_package(f.source)
        f.nodes=region_nodes(f.trees)
        f.report=inventory(f.source,f.trees,f.raw,f.entries)

    def test_revise_nonempty_identity_prose_and_filename(self):
        self.populated_fixture(); f=self.f
        find=lambda text: next(k for k,n in f.nodes.items() if n.tag==W+'p' and text_of(n)==text)
        process=next(t for t in f.report['tables'] if t['parent_cell'])['rows'][0][1]
        result=self.run_patch({'mode':'revise','filename_replacements':[
            {'old':'99','field':'student_id_last2'}, {'old':'旧同学甲','field':'name'}, {'old':'旧测试班','field':'class'}],
            'operations':[
                {'op':'replace','target':find('学生：旧同学甲'),'old':'旧同学甲','field':'name'},
                {'op':'replace','target':find('班级：旧测试班'),'old':'旧测试班','field':'class'},
                {'op':'replace','target':find('学号：000099'),'old':'000099','field':'student_id'},
                {'target':process,'expect':'原步骤错误内容','text':'按手册先检查连接，再执行规定步骤。'},
                {'target':f.cell('心得'),'expect':'原心得第一段\n原心得第二段','field':'reflection','text':'本次练习明确了操作顺序。\n应按要求记录观察情况。','font_size_pt':12},
                {'target':f.cell('教师评语'),'expect':'原普通评语','field':'teacher_comment','text':'步骤表述清楚，注意事项与本次任务对应。','font_size_pt':12}]})
        self.assertEqual(result['mode'],'revise')
        self.assertEqual(Path(result['output']).name,'07测试同学测试班实训.docx')
        _, entries, trees=read_package(result['output'])
        body=text_of(trees['word/document.xml'])
        self.assertIn('学生：测试同学',body)
        self.assertIn('文献作者：旧同学甲；指导教师：教师乙',body)
        self.assertIn('教师实际批阅：原记录保留',body)
        self.assertNotIn('原心得',body); self.assertNotIn('原普通评语',body)
        self.assertNotIn('原步骤错误内容',body)
        self.assertIn('班级：测试班',text_of(trees['word/header1.xml']))
        self.assertIn('学号：000007',text_of(trees['word/footer1.xml']))
        self.assertEqual(len(trees['word/document.xml'].xpath('.//w:ind[@w:firstLineChars="200"]',namespaces=NS)),3)
        def skeleton(tree):
            cloned=copy.deepcopy(tree)
            for p in list(cloned.iter(W+'p')): p.getparent().remove(p)
            return ET.tostring(cloned,method='c14n')
        self.assertEqual(skeleton(trees['word/document.xml']),skeleton(f.trees['word/document.xml']))
        before={i.filename:data for i,data in f.entries}
        for info,data in entries:
            if info.filename not in ('word/document.xml','word/header1.xml','word/footer1.xml'):
                self.assertEqual(data,before[info.filename],info.filename)

    def test_revise_keeps_nonempty_precondition_and_image_guards(self):
        self.populated_fixture(); f=self.f
        result=self.run_patch({'mode':'revise','operations':[{'target':f.cell('教师评语'),'field':'teacher_comment','text':'新评语'}]},success=False)
        self.assertIn('precondition failed',result['error'])
        result=self.run_patch({'mode':'revise','operations':[{'target':f.cell('指导教师'),'field':'teacher','text':'教师乙'}]})
        self.assertEqual(result['skipped'][0]['reason'],'teacher_signature_image_preserved')

    def test_filename_mapping_is_unique_and_not_cascaded(self):
        profile={'name':'旧乙','class':'新班','student_id':'000007'}
        self.assertEqual(output_name(Path('旧甲+旧乙.docx'),profile,[{'old':'旧甲','field':'name'},{'old':'旧乙','field':'class'}]),'旧乙新班.docx')
        cases=[('旧甲旧甲.docx',[{'old':'旧甲','field':'name'}]),
               ('旧甲.docx',[{'old':'不存在','field':'name'}]),
               ('旧同学甲.docx',[{'old':'旧同学甲','field':'name'},{'old':'同学甲','field':'class'}]),
               ('aaa.docx',[{'old':'aa','field':'name'}]),
               ('旧甲.docx',[{'old':'旧甲','field':'teacher'}])]
        for name,mapping in cases:
            with self.subTest(name=name,mapping=mapping):
                with self.assertRaises(ValueError): output_name(Path(name),profile,mapping)

    def test_mode_restrictions_and_original_protection(self):
        f=self.f
        for patch in ({'mode':'unknown'},{'filename_replacements':[{'old':'x','field':'name'}]}, {'mode':'revise','filename_replacements':[]}):
            self.run_patch(patch,success=False)
        self.run_patch({'mode':'revise','filename_replacements':[{'old':'x','field':'name'}]},output=f.root/'explicit.docx',success=False)
        source=f.root/'原件.docx'; source.write_bytes(f.raw); f.source=source
        result=self.run_patch({'mode':'revise'},output=source,success=False)
        self.assertIn('original',result['error'])
        unknown=f.root/'未知.docx'; unknown.write_bytes(b'unknown user file')
        self.run_patch({'mode':'revise'},output=unknown,success=False)
        self.assertEqual(unknown.read_bytes(),b'unknown user file')


if __name__=='__main__': unittest.main()
